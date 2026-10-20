"""CAS+ as a table a person can read, in a spreadsheet.

A CAS+ document is mostly the breakdown of every descriptor over every cell and
the provenance of every label it inherited. Both run to thousands of entries and
neither says what a cell type is called, where it sits, or what it was mapped to.
Strip those two out and what is left — identity, hierarchy, size, the ontology
mapping, the evidence behind it — is a few hundred short lines.

So this carries every annotation field that has an honest tabular form, and
leaves out only the three that do not: ``composition`` and
``transferred_annotations``, each of which is a table in its own right, and
``author_annotation_fields``, an open-ended object whose keys differ per atlas.
``minimal`` drops the free-text columns, for a table to be scanned rather than
read.

Cell sets are identified by accession. A label is not unique — an atlas may use
one at two levels of its hierarchy, for cell sets with different members and
different terms — so a selection may legitimately return several, and the table
carries the accession to tell them apart.

Nothing here is computed that the document does not already state, and nothing
here calls a model.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class Column:
    """A column: where it comes from, and how much room it needs.

    The width and the wrapping live here rather than in the spreadsheet writer
    because how much room a column needs is a fact about what it holds.
    """

    name: str
    field: str
    width: int
    wrap: bool = False
    free_text: bool = False


#: Every column, in the order written: identity, hierarchy, size, the mapping,
#: the evidence behind it, then prose.
COLUMNS: tuple[Column, ...] = (
    Column("accession", "cell_set_accession", 20),
    Column("labelset", "labelset", 9),
    Column("label", "cell_label", 30),
    Column("full name", "cell_fullname", 38),
    Column("parent", "parent_cell_set_accession", 30),
    Column("cells", "n_cells", 10),
    Column("cl_id", "cell_ontology_term_id", 14),
    Column("cl_term", "cell_ontology_term", 32),
    Column("cl_match_type", "cell_ontology_match_type", 14),
    Column("markers", "marker_gene_evidence", 30, wrap=True),
    Column("negative_markers", "negative_marker_gene_evidence", 24, wrap=True),
    Column("synonyms", "synonyms", 28, wrap=True),
    Column("rationale_dois", "rationale_dois", 26, wrap=True),
    Column("cl_mapping_notes", "cell_ontology_mapping_notes", 50, wrap=True, free_text=True),
    Column("comment", "comment", 60, wrap=True, free_text=True),
    Column("rationale", "rationale", 60, wrap=True, free_text=True),
)

BY_NAME = {c.name: c for c in COLUMNS}

#: What :func:`project_paths.outline` emitted before this module existed, whose
#: callers read its columns by position.
LEGACY_OUTLINE = ("labelset", "label", "full name", "parent", "cells")

#: Separator inside a list-valued cell, matching the existing convention.
LIST_SEPARATOR = ";"

#: Written at the head of a text file so a spreadsheet reads it as UTF-8 rather
#: than as the machine's legacy codepage.
BOM = "﻿"

FORMATS = ("xlsx", "tsv", "csv")


class CasReportError(RuntimeError):
    """Raised when a selection cannot be resolved, or a format cannot be written."""


def chosen(*, minimal: bool = False, only: list[str] | None = None) -> tuple[Column, ...]:
    """The columns to write.

    Args:
        minimal: drop the free-text columns.
        only: exactly these column names, in this order.

    Returns:
        The columns, in the order they will be written.

    Raises:
        CasReportError: ``only`` names a column that does not exist.
    """
    if only is not None:
        unknown = [n for n in only if n not in BY_NAME]
        if unknown:
            raise CasReportError(
                f"no such column: {', '.join(unknown)}. "
                f"available: {', '.join(c.name for c in COLUMNS)}"
            )
        return tuple(BY_NAME[n] for n in only)
    return tuple(c for c in COLUMNS if not (minimal and c.free_text))


def _searchable(annotation: dict[str, Any]) -> str:
    parts = [annotation.get("cell_label", ""), annotation.get("cell_fullname") or ""]
    parts.extend(annotation.get("synonyms") or [])
    return " ".join(parts).lower()


def _descendants(root: str, by_parent: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    stack = list(by_parent.get(root, []))
    while stack:
        node = stack.pop()
        out.append(node)
        stack.extend(by_parent.get(node.get("cell_set_accession") or "", []))
    return out


def select(
    cas: dict[str, Any],
    *,
    labelset: str | None = None,
    match: str | None = None,
    under: str | None = None,
    accessions: list[str] | None = None,
) -> list[dict[str, Any]]:
    """The annotations a request names, in document order.

    Filters combine: each narrows what the ones before it left. ``under`` naming
    a label used at two levels takes what sits under both, because the result is
    one table keyed by accession and the two subtrees are distinguishable in it.

    Args:
        cas: the CAS+ document.
        labelset: keep only this level.
        match: keep annotations whose label, full name or synonyms contain this.
        under: keep everything beneath the annotation(s) with this label.
        accessions: keep exactly these cell sets.

    Returns:
        The annotations selected, in document order.

    Raises:
        CasReportError: an accession or an ``under`` label is not in the
            document. Silently returning a smaller table would read as a finding.
    """
    annotations = cas.get("annotations") or []
    by_parent: dict[str, list[dict[str, Any]]] = {}
    for a in annotations:
        by_parent.setdefault(a.get("parent_cell_set_accession") or "", []).append(a)

    kept = list(annotations)
    if accessions:
        known = {a.get("cell_set_accession") for a in annotations}
        missing = [x for x in accessions if x not in known]
        if missing:
            raise CasReportError(f"not in the document: {', '.join(sorted(missing))}")
        wanted = set(accessions)
        kept = [a for a in kept if a.get("cell_set_accession") in wanted]
    if under:
        roots = [a for a in annotations if a.get("cell_label") == under]
        if not roots:
            raise CasReportError(f"not in the document: {under}")
        beneath = {
            id(a)
            for root in roots
            for a in _descendants(root.get("cell_set_accession") or "", by_parent)
        }
        kept = [a for a in kept if id(a) in beneath]
    if labelset:
        kept = [a for a in kept if a.get("labelset") == labelset]
    if match:
        needle = match.lower()
        kept = [a for a in kept if needle in _searchable(a)]
    return kept


def _value(annotation: dict[str, Any], column: Column, label_of: dict[str, str]) -> str:
    """One cell, as text.

    A list is joined rather than reduced to its first element, so an ambiguous
    ontology mapping reads as ambiguous rather than as a decision. The parent is
    resolved to its label, because an accession says nothing to a reader.
    """
    if column.field == "parent_cell_set_accession":
        return label_of.get(annotation.get("parent_cell_set_accession") or "", "")
    raw = annotation.get(column.field)
    if column.field == "cell_fullname" and raw == annotation.get("cell_label"):
        return ""  # a full name identical to the label is noise in every row
    if raw is None:
        return ""
    if isinstance(raw, list):
        return LIST_SEPARATOR.join(str(x) for x in raw)
    return str(raw)


def rows(
    cas: dict[str, Any],
    columns: tuple[Column, ...],
    *,
    labelset: str | None = None,
    match: str | None = None,
    under: str | None = None,
    accessions: list[str] | None = None,
) -> list[list[str]]:
    """The selected annotations as rows of text, matching ``columns``."""
    label_of = {
        a["cell_set_accession"]: a["cell_label"]
        for a in (cas.get("annotations") or [])
        if a.get("cell_set_accession")
    }
    selected = select(cas, labelset=labelset, match=match, under=under, accessions=accessions)
    return [[_value(a, c, label_of) for c in columns] for a in selected]


def _flatten(text: str) -> str:
    """Collapse the tabs and newlines a delimited row cannot carry."""
    return " ".join(text.split()) if ("\t" in text or "\n" in text or "\r" in text) else text


def to_text(columns: tuple[Column, ...], table: list[list[str]], *, delimiter: str = "\t") -> str:
    """The table as delimited text.

    Tabs and newlines inside a field are collapsed rather than quoted: a free-text
    column carrying a newline would otherwise split one annotation across several
    rows, which reads as extra cell types.
    """
    if delimiter == ",":
        buffer = io.StringIO()
        writer = csv.writer(buffer, lineterminator="\n")
        writer.writerow([c.name for c in columns])
        writer.writerows([[_flatten(v) for v in row] for row in table])
        return buffer.getvalue().rstrip("\n")
    lines = [delimiter.join(c.name for c in columns)]
    lines += [delimiter.join(_flatten(v) for v in row) for row in table]
    return "\n".join(lines)


def write_text(
    columns: tuple[Column, ...],
    table: list[list[str]],
    out: Path,
    *,
    delimiter: str = "\t",
    bom: bool = True,
) -> None:
    """Write delimited text, with the byte-order mark a spreadsheet needs."""
    body = to_text(columns, table, delimiter=delimiter)
    out.write_text((BOM if bom else "") + body + "\n", encoding="utf-8")


def write_xlsx(columns: tuple[Column, ...], table: list[list[str]], out: Path) -> None:
    """Write a spreadsheet laid out to be read, rather than reformatted on arrival.

    The header is frozen and filterable, each column gets the width its content
    needs, and the columns holding more than a few words wrap instead of running
    out under their neighbours.

    Raises:
        CasReportError: openpyxl is not installed.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - an install problem, not a code path
        raise CasReportError(
            "writing xlsx needs openpyxl; install it, or ask for --format tsv"
        ) from exc

    book = Workbook()
    sheet = book.active
    sheet.title = "annotations"
    sheet.append([c.name for c in columns])
    for row in table:
        sheet.append(row)

    wrapped = Alignment(vertical="top", wrap_text=True)
    top = Alignment(vertical="top")
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = column.width
        header = sheet.cell(row=1, column=index)
        header.font = Font(bold=True)
        header.alignment = top
        if column.wrap:
            for row_index in range(2, len(table) + 2):
                sheet.cell(row=row_index, column=index).alignment = wrapped

    sheet.freeze_panes = "A2"
    if table:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(table) + 1}"
    book.save(out)


def annotations_table(
    cas: dict[str, Any],
    *,
    labelset: str | None = None,
    match: str | None = None,
    under: str | None = None,
    accessions: list[str] | None = None,
    minimal: bool = False,
    only: list[str] | None = None,
    delimiter: str = "\t",
) -> str:
    """The annotation hierarchy and its mappings, as delimited text.

    Args:
        cas: the CAS+ document.
        labelset: keep only this level.
        match: keep annotations whose label, full name or synonyms contain this.
        under: keep everything beneath the annotation(s) with this label.
        accessions: keep exactly these cell sets.
        minimal: drop the free-text columns.
        only: exactly these column names, in this order.
        delimiter: tab, or a comma for CSV.

    Returns:
        A header line and one line per annotation, in document order. Empty
        where a field is absent, so the columns stay aligned.

    Raises:
        CasReportError: the selection names something not in the document, or
            ``only`` names a column that does not exist.
    """
    columns = chosen(minimal=minimal, only=only)
    table = rows(cas, columns, labelset=labelset, match=match, under=under, accessions=accessions)
    return to_text(columns, table, delimiter=delimiter)


# ------------------------------------------------------------------
# CLI
# ------------------------------------------------------------------


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="python -m atlas_chat.cli_report",
        description="CAS+ as a table a person can read.",
    )
    sub = parser.add_subparsers(dest="command", required=True)

    table = sub.add_parser("annotations", help="the annotation hierarchy and its mappings")
    table.add_argument("--cas", required=True, help="the CAS+ document")
    table.add_argument("--labelset", help="keep only this level")
    table.add_argument("--match", help="keep annotations whose label, full name or synonyms match")
    table.add_argument("--under", help="keep everything beneath this annotation")
    table.add_argument(
        "--accession",
        action="append",
        help="repeatable; names a cell set exactly, for a label used at more than one level",
    )
    table.add_argument(
        "--minimal",
        action="store_true",
        help="drop the free-text columns: " + ", ".join(c.name for c in COLUMNS if c.free_text),
    )
    table.add_argument(
        "--columns",
        help="comma-separated column names, in the order wanted; overrides --minimal",
    )
    table.add_argument(
        "--format",
        choices=FORMATS,
        default="xlsx",
        help="default xlsx, laid out for reading in a spreadsheet",
    )
    table.add_argument("--out", help="write here; required for xlsx, which is not text")
    table.add_argument(
        "--no-bom",
        action="store_true",
        help="omit the byte-order mark from tsv/csv; without it a spreadsheet reads "
        "the file in the machine's legacy codepage",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    cas = json.loads(Path(args.cas).read_text(encoding="utf-8"))
    only = [c.strip() for c in args.columns.split(",")] if args.columns else None
    delimiter = "," if args.format == "csv" else "\t"
    try:
        columns = chosen(minimal=args.minimal, only=only)
        table = rows(
            cas,
            columns,
            labelset=args.labelset,
            match=args.match,
            under=args.under,
            accessions=args.accession,
        )
        if args.format == "xlsx":
            if not args.out:
                raise CasReportError("xlsx is not text: give --out, or ask for --format tsv")
            write_xlsx(columns, table, Path(args.out))
        elif args.out:
            write_text(columns, table, Path(args.out), delimiter=delimiter, bom=not args.no_bom)
        else:
            print(to_text(columns, table, delimiter=delimiter))
            return 0
    except CasReportError as exc:
        print(str(exc))
        return 2
    print(f"{len(table)} annotations, {len(columns)} columns -> {args.out}")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
